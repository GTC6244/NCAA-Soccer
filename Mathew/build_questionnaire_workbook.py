#!/usr/bin/env python3
"""Build the men's-soccer recruiting questionnaire workbook for Mathew Paul.
Catalogs the data each school's 'register your interest' form requires and
fills in Mathew's answers from the supplied profile. Does NOT submit anything.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ---------- styling helpers ----------
HDR = Font(bold=True, color="FFFFFF", size=11)
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
GRP_FILL = PatternFill("solid", fgColor="D9E1F2")
GRP_FONT = Font(bold=True, color="1F4E79")
NEED_FILL = PatternFill("solid", fgColor="FCE4D6")   # orange = still needed
ANS_FILL = PatternFill("solid", fgColor="E2EFDA")    # green-ish for answer col
REQ_FONT = Font(bold=True, color="C00000")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CTR = Alignment(horizontal="center", vertical="center")

NEED = "⚠ need from Mathew"   # warning marker

# =====================================================================
# SHEET 1 — Mathew profile
# =====================================================================
ws = wb.active
ws.title = "Mathew Profile"
profile = [
    ("PLAYER", ""),
    ("Full name", "Mathew Paul"),
    ("First / Last", "Mathew / Paul"),
    ("Date of birth", "June 5, 2009"),
    ("City/State of birth", "Kitchener, Ontario, Canada"),
    ("Player email", "Mathewp@axom.ca"),
    ("Mobile phone", "+1 (226) 507-3758"),
    ("Home phone", "(519) 893-2340"),
    ("Home address", "183 Pine Valley Drive, Kitchener, Ontario"),
    ("City", "Kitchener"),
    ("Province/State", "Ontario"),
    ("Postal code", "N2P 2V8"),
    ("Country", "Canada"),
    ("Hometown", "Kitchener, Ontario, Canada"),
    ("Nationality / Citizenship", "Canadian"),
    ("Languages", "English & French (writes/speaks French: Yes)"),
    ("Instagram", "Mathewpaul009"),
    ("Twitter/X", "— none provided"),
    ("Highlight video", "https://go.axom.ca/mathew-fb"),
    ("ACADEMIC", ""),
    ("Graduation year", "2027 (corrected — '2009' supplied is birth year; confirm)"),
    ("GPA", "4.0"),
    ("Core-course GPA", "4.0"),
    ("Class rank", NEED),
    ("SAT", NEED + " / or N/A"),
    ("ACT", NEED + " / or N/A"),
    ("TOEFL", "N/A (fluent English)"),
    ("Intended major", "Math, Engineering, Computer Science"),
    ("High school", "Huron Heights Secondary School (Kitchener, ON)"),
    ("Transcript", "to upload"),
    ("NCAA Eligibility Center registered?", "Yes"),
    ("NCAA ID", NEED),
    ("ATHLETIC", ""),
    ("Primary sport", "Men's Soccer"),
    ("Primary position", "Centre Back (CB)"),
    ("Secondary position", "Wing Back (WB)"),
    ("Dominant foot", "Right"),
    ("Height", "6'2\" (188 cm)"),
    ("Weight", "179 lb (81 kg)"),
    ("Jersey #", "4"),
    ("Current club / team", "Waterloo United – BVB IA Waterloo, U17 OPDL (#4)"),
    ("Club location", "Waterloo, Ontario, Canada"),
    ("Club coach (name/phone/email)", "Ricky Gomes, ricky.gomes@waterloounited.com"),
    ("Interested in College / University / both", "Both"),
    ("Interested in US schools", "Yes"),
    ("Willing to move away from home", "Yes"),
    ("ACHIEVEMENTS / HONOURABLE MENTIONS", ""),
    ("Soccer record", (
        "2024 BVB IA Waterloo – U15 WRSL (Team Captain); U16 6x full games; U18 tournaments. "
        "2024 Germany – 1 wk w/ Hamburg SV. 2024 Spain – Rayo Alcobendas (League of Honour club) 2.5 months. "
        "2024 BVB IA East Regional Team (US/CAN). 2024 BVB IA North American National Team (only Canadian member). "
        "2025 BVB IA School of Excellence – invited to Germany 1.5 wks (costs covered by BVB Germany); "
        "TV appearance @ BVB Bundesliga game. 2025 Portugal 2 wks training & Iber Cup play. "
        "2025 BVB IA Waterloo – U16 OPDL.")),
    ("FAMILY", ""),
    ("Parent/guardian name", "Brendon Paul"),
    ("Parent email", "brendonp@ixiomsoftware.com"),
    ("Siblings", "Isabelle (sister)"),
]
ws["A1"] = "Mathew Paul — Recruiting Profile (source data for all questionnaires)"
ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
ws.merge_cells("A1:B1")
r = 3
ws.cell(r, 1, "Field").font = HDR; ws.cell(r, 1).fill = HDR_FILL
ws.cell(r, 2, "Value").font = HDR; ws.cell(r, 2).fill = HDR_FILL
r += 1
for f, v in profile:
    if v == "" and f.isupper():
        c = ws.cell(r, 1, f); c.font = GRP_FONT; c.fill = GRP_FILL
        ws.cell(r, 2, "").fill = GRP_FILL
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.cell(r,1).value = f
    else:
        ws.cell(r, 1, f).font = Font(bold=True)
        vc = ws.cell(r, 2, v)
        vc.alignment = WRAP
        if isinstance(v, str) and NEED in v:
            vc.fill = NEED_FILL
    for col in (1, 2):
        ws.cell(r, col).border = BORDER
    r += 1
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 90

# =====================================================================
# SHEET 2 — Questionnaire directory (all 28 schools)
# =====================================================================
wd = wb.create_sheet("Questionnaire Directory")
dir_hdr = ["School", "Tier", "Division", "Platform", "Form URL", "Status", "Notes"]
directory = [
 ("UCF (Central Florida)","Core","NCAA D1","SIDEARM (custom bio form)","https://ucfknights.com/ucf-mens-soccer-student-athlete-questionnaire","Form open","Essay/bio-heavy; NO email/phone/GPA/test fields on the form."),
 ("USF (South Florida)","Core","NCAA D1","JumpForward","https://college.jumpforward.com/questionnaire.aspx?iid=1625&sportid=21","Form open","Full recruiting questionnaire incl. parents & siblings."),
 ("South Carolina","Core","NCAA D1","JumpForward","https://college.jumpforward.com/questionnaire.aspx?iid=404&sportid=21","Form open","Required markers visible; HUDL/YouTube fields; no family section."),
 ("Ohio State","Core","NCAA D1","FieldLevel","https://www.fieldlevel.com/7cwan776/soccermen/recruiting","Account required","Must create a FieldLevel profile. Fields = standard FieldLevel profile (generic)."),
 ("Penn State","Core","NCAA D1","ARMS / Teamworks","https://questionnaires.armssoftware.com/arms/public/questionnaire/0e60bc2ee1c0","Form open","Fields per ARMS standard template (not individually verified)."),
 ("Boston University","Core","NCAA D1","ARMS / Teamworks","https://questionnaires.armssoftware.com/3585708e2ece","Form open","Fields per ARMS standard template (not individually verified)."),
 ("Northeastern","Core","NCAA D1","ARMS / Teamworks","https://questionnaires.armssoftware.com/e2e3db7f1062","Form open","Verified ARMS field set. Gate: First/Last/Email."),
 ("UMass Lowell","Core","NCAA D1","ARMS / Teamworks","https://my.armssoftware.com/arms/public/questionnaire/34ce2eed1194","Form open","'UML Men's Soccer'. Standard ARMS template."),
 ("IU Indianapolis","Core","NCAA D1","ARMS / Teamworks","https://questionnaires.armssoftware.com/8792f61a8ff2","Form open","Standard ARMS template (single mobile phone). Recruiting guide PDF also on site."),
 ("SMU","Core","NCAA D1","— (none posted)","—","Email coach","No online questionnaire. Email HC Kevin Hudson: khudson@smu.edu."),
 ("UC Santa Barbara","Core","NCAA D1","Exposure microsite","https://ucsbmenssoccer.exposure.co","Blocked (WAF)","Recruiting page blocks all automated access; open in a normal browser."),
 ("NJIT","Core","NCAA D1","— (no soccer form)","https://njithighlanders.com/documents/2015/7/21/NJIT_prospect_info_form.pdf","No soccer form","Only an outdated generic PDF (labeled tennis). Email HC: barboto@njit.edu."),
 ("FIU (Florida Intl)","Extra","NCAA D1","JumpForward","https://college.jumpforward.com/questionnaire.aspx?iid=376&sportid=21","Form open","Full questionnaire incl. counselor, parents & siblings."),
 ("NC State","Extra","NCAA D1","ARMS / Teamworks","https://questionnaires.armssoftware.com/936d7daa3fc8","Form open","Longest form – many extra academic/personal/interest questions."),
 ("Georgia Southern","Extra","NCAA D1","ARMS / Teamworks","https://questionnaires.armssoftware.com/350ec954fad8","Form open","Standard ARMS + dominant foot, guidance counselor."),
 ("Georgia State","Extra","NCAA D1","— (none posted)","—","Email coach","No online form posted – email direct."),
 ("Kennesaw State","Extra","NCAA D1","— (not found)","—","Phone athletics","No form/contact retrievable. Call athletics dept."),
 ("UNC Chapel Hill","Extra","NCAA D1","— (none posted)","—","Email coach","No online form posted."),
 ("UNC Charlotte","Extra","NCAA D1","— (not retrieved)","—","Needs sweep","Coach-contact / form not yet retrieved."),
 ("UNC Wilmington","Extra","NCAA D1","— (not retrieved)","—","Needs sweep","Coach-contact / form not yet retrieved."),
 ("UNC Greensboro","Extra","NCAA D1","— (not retrieved)","—","Needs sweep","Coach-contact / form not yet retrieved."),
 ("Florida Atlantic","Extra","NCAA D1","— (not retrieved)","—","Needs sweep","Coach-contact / form not yet retrieved."),
 ("Clemson","Extra","NCAA D1","— (not retrieved)","—","Needs sweep","Coach-contact / form not yet retrieved."),
 ("Nova Southeastern","Extra","NCAA D2","— (not retrieved)","—","Needs sweep","Coach-contact / form not yet retrieved."),
 ("Keiser","Extra","NAIA","— (not retrieved)","—","Needs sweep","Coach-contact / form not yet retrieved."),
 ("University of Toronto","Canada","U SPORTS","SIDEARM prospect form","https://varsityblues.ca/sb_output.aspx?form=3","Form open","Shared all-sports 'Prospective Athlete Form'; pick Men's Soccer."),
 ("York University","Canada","U SPORTS","SIDEARM prospect form","https://yorkulions.ca/sb_output.aspx?form=3","Form open","Shared all-sports form; sport dropdown includes Men's Soccer."),
 ("Western University","Canada","U SPORTS","— (not posted)","—","Email coach","No online form posted – email direct."),
]
for j, h in enumerate(dir_hdr, 1):
    c = wd.cell(1, j, h); c.font = HDR; c.fill = HDR_FILL; c.border = BORDER; c.alignment = CTR
for i, row in enumerate(directory, 2):
    for j, val in enumerate(row, 1):
        c = wd.cell(i, j, val); c.border = BORDER; c.alignment = WRAP
        if j == 6 and val in ("Blocked (WAF)","No soccer form","Needs sweep","Account required"):
            c.fill = NEED_FILL
widths = [24, 8, 11, 24, 58, 15, 46]
for j, w in enumerate(widths, 1):
    wd.column_dimensions[get_column_letter(j)].width = w
wd.freeze_panes = "A2"

# =====================================================================
# SHEET 3 — Fields x School matrix with Mathew's answers
# =====================================================================
mx = wb.create_sheet("Fields x School (with answers)")

# ordered field rows: (key, group, label, mathew_answer)
A = {
 "p_full":"Mathew Paul","p_first":"Mathew","p_middle":"","p_last":"Paul","p_pref":"",
 "p_pron":"","p_photo":"to provide","p_dob":"June 5, 2009","p_bplace":"Kitchener, ON, Canada","p_email":"Mathewp@axom.ca",
 "p_hphone":"(519) 893-2340","p_mphone":"+1 (226) 507-3758","p_addr":"183 Pine Valley Drive, Kitchener, Ontario","p_city":"Kitchener","p_state":"Ontario",
 "p_zip":"N2P 2V8","p_country":"Canada","p_home":"Kitchener, ON, Canada","p_citz":"Canadian",
 "p_ig":"Mathewpaul009","p_tw":"— none","p_fb":"— none","p_hob":"","p_lang":"English & French",
 "a_grad":"2027 (corrected; confirm)","a_incls":"Freshman / 2027","a_hs":"Huron Heights Secondary School","a_hsaddr":"Kitchener, Ontario","a_hsph":NEED,"a_couns":NEED,
 "a_gpa":"4.0","a_cgpa":"4.0","a_rank":NEED,"a_sat":NEED,"a_act":NEED,"a_toefl":"N/A (fluent English)",
 "a_major":"Math, Engineering, CS","a_career":"","a_trans":"to upload","a_testf":"to upload",
 "a_ncaa":"Yes","a_ncaaid":NEED,"a_applied":"No","a_hon":NEED+" / optional",
 "a_other":"(optional – target list)","a_colprev":"N/A (in high school)",
 "t_ppos":"Centre Back (CB)","t_spos":"Wing Back (WB)","t_sport":"Men's Soccer","t_ht":"6'2\" (188 cm)",
 "t_wt":"179 lb (81 kg)","t_foot":"Right","t_jersey":"4",
 "t_club":"Waterloo United / BVB IA Waterloo U17 OPDL","t_cloc":"Waterloo, ON, Canada",
 "t_ccoach":"Ricky Gomes, ricky.gomes@waterloounited.com","t_cpos":"CB, #4","t_tourn":"Iber Cup 2025 (Portugal); Germany 2024; Spain 2024",
 "t_vid":"https://go.axom.ca/mathew-fb","t_hudl":"— none","t_stats":"See achievements (BVB NA Nat'l Team, School of Excellence, U15 Captain, etc.)",
 "t_ppg":"N/A (defender)","t_osport":"",
 "f_p1":"Brendon Paul — brendonp@ixiomsoftware.com","f_p2":"","f_pemail":"brendonp@ixiomsoftware.com",
 "f_pocc":"","f_sib":"Isabelle (sister)","f_coach":"Ricky Gomes, ricky.gomes@waterloounited.com","f_infl":"","f_relath":"","f_relsch":"","f_live":"",
 "o_why":"to write per school","o_vol":"","o_bio":"See achievements","o_rate":"per school",
 "o_prio":"","o_visit":"","o_free":"See achievements","o_usid":NEED+" / if applicable","o_finaid":NEED+" (confirm)",
}
ROWS = [
 ("PERSONAL",None,None),
 ("p_full","Full name","p_full"),("p_first","First name","p_first"),("p_middle","Middle name","p_middle"),
 ("p_last","Last name","p_last"),("p_pref","Preferred name","p_pref"),("p_pron","Name pronunciation","p_pron"),
 ("p_photo","Headshot / photo","p_photo"),("p_dob","Date of birth","p_dob"),("p_bplace","City/State of birth","p_bplace"),
 ("p_email","Email (player)","p_email"),("p_hphone","Home phone","p_hphone"),("p_mphone","Mobile phone","p_mphone"),
 ("p_addr","Street address","p_addr"),("p_city","City","p_city"),("p_state","State/Province","p_state"),
 ("p_zip","Zip/Postal","p_zip"),("p_country","Country","p_country"),("p_home","Hometown","p_home"),
 ("p_citz","Nationality/Citizenship","p_citz"),("p_ig","Instagram","p_ig"),("p_tw","Twitter/X","p_tw"),
 ("p_fb","Facebook","p_fb"),("p_hob","Hobbies","p_hob"),("p_lang","Languages","p_lang"),
 ("ACADEMIC",None,None),
 ("a_grad","Graduation year","a_grad"),("a_incls","Incoming class (Fr/So/Jr/Sr)","a_incls"),
 ("a_hs","High school name","a_hs"),("a_hsaddr","HS address/city/state","a_hsaddr"),
 ("a_hsph","HS phone/website/CEEB","a_hsph"),("a_couns","Guidance counselor","a_couns"),
 ("a_gpa","GPA","a_gpa"),("a_cgpa","Core-course GPA","a_cgpa"),("a_rank","Class rank","a_rank"),
 ("a_sat","SAT (+subscores/date)","a_sat"),("a_act","ACT (+subscores/date)","a_act"),("a_toefl","TOEFL","a_toefl"),
 ("a_major","Intended major","a_major"),("a_career","Career interest","a_career"),
 ("a_trans","Transcript (upload)","a_trans"),("a_testf","Test scores (upload)","a_testf"),
 ("a_ncaa","NCAA Eligibility Ctr registered?","a_ncaa"),("a_ncaaid","NCAA ID","a_ncaaid"),
 ("a_applied","Applied to this school?","a_applied"),("a_hon","Academic honors/clubs","a_hon"),
 ("a_other","Other schools of interest","a_other"),("a_colprev","Previous college / college GPA","a_colprev"),
 ("ATHLETIC",None,None),
 ("t_ppos","Primary position","t_ppos"),("t_spos","Secondary position","t_spos"),
 ("t_sport","Primary sport (dropdown)","t_sport"),("t_ht","Height","t_ht"),("t_wt","Weight","t_wt"),
 ("t_foot","Dominant foot","t_foot"),("t_jersey","Jersey #","t_jersey"),
 ("t_club","Club / current team","t_club"),("t_cloc","Club city/state","t_cloc"),
 ("t_ccoach","Club coach (name/phone/email)","t_ccoach"),("t_cpos","Club position / jersey #","t_cpos"),
 ("t_tourn","Tournaments","t_tourn"),("t_vid","Highlight / recruiting video","t_vid"),
 ("t_hudl","HUDL / YouTube","t_hudl"),("t_stats","Stats / athletic honors","t_stats"),
 ("t_ppg","Points/Assists per game","t_ppg"),("t_osport","Other sports played","t_osport"),
 ("FAMILY",None,None),
 ("f_p1","Parent/Guardian 1 (name & contact)","f_p1"),("f_p2","Parent/Guardian 2","f_p2"),
 ("f_pemail","Parent email","f_pemail"),("f_pocc","Parent occupation/employer","f_pocc"),
 ("f_sib","Siblings (names/ages)","f_sib"),("f_coach","Coaches (related contact)","f_coach"),
 ("f_infl","Most influential people","f_infl"),("f_relath","Notable athlete relatives","f_relath"),
 ("f_relsch","Relatives who attended school","f_relsch"),("f_live","Who you live with / parents' marital status","f_live"),
 ("OTHER",None,None),
 ("o_why","Why this school?","o_why"),("o_vol","Volunteer activities","o_vol"),
 ("o_bio","Additional bio / 'something we should know'","o_bio"),("o_rate","Rate interest level","o_rate"),
 ("o_prio","Top priorities choosing a program","o_prio"),("o_visit","Schools you plan to visit","o_visit"),
 ("o_free","Free-text / additional info","o_free"),("o_usid","U SPORTS ID","o_usid"),
 ("o_finaid","Financial aid intent","o_finaid"),
]

R_,O_ = "R","●"
# ---- ARMS standard template (Northeastern etc.) ----
arms = {"p_first":R_,"p_last":R_,"p_email":R_,"a_grad":R_,
 "p_photo":O_,"p_hphone":O_,"p_mphone":O_,"p_addr":O_,"p_city":O_,"p_state":O_,"p_zip":O_,"p_country":O_,
 "p_dob":O_,"p_fb":O_,"p_tw":O_,"p_hob":O_,
 "f_p1":O_,"f_coach":O_,"a_couns":O_,"f_infl":O_,"f_sib":O_,
 "a_hs":O_,"a_trans":O_,"a_testf":O_,"a_gpa":O_,"a_rank":O_,"a_major":O_,"a_ncaa":O_,"a_ncaaid":O_,
 "a_sat":O_,"a_act":O_,"a_hon":O_,"a_other":O_,
 "t_ppos":O_,"t_ht":O_,"t_wt":O_,"t_foot":O_,"t_vid":O_,"t_club":O_,"t_cpos":O_,"t_tourn":O_,"t_osport":O_,"t_stats":O_}
northeastern = dict(arms)
umasslowell = dict(arms)
iuindy = dict(arms); iuindy.pop("p_hphone")
pennstate = dict(arms)
bostonu = dict(arms)

ncstate = {"a_grad":R_,"p_first":R_,"p_last":R_,"p_email":R_,"p_mphone":R_,"a_ncaa":R_,"a_applied":R_,
 "a_hs":R_,"a_hsaddr":R_,
 "p_photo":O_,"p_middle":O_,"p_pref":O_,"p_addr":O_,"p_city":O_,"p_state":O_,"p_zip":O_,"p_country":O_,
 "p_dob":O_,"p_hphone":O_,"a_ncaaid":O_,"p_fb":O_,"p_tw":O_,"p_hob":O_,
 "f_p1":O_,"f_coach":O_,"f_infl":O_,"f_pocc":O_,"f_live":O_,"f_sib":O_,"f_relath":O_,"o_bio":O_,
 "t_ppos":O_,"t_ppg":O_,"t_ht":O_,"t_wt":O_,"t_jersey":O_,"t_osport":O_,"t_vid":O_,"t_club":O_,"t_ccoach":O_,"t_stats":O_,
 "a_gpa":O_,"a_cgpa":O_,"a_rank":O_,"a_act":O_,"a_sat":O_,"a_major":O_,"a_career":O_,"a_hsph":O_,"a_trans":O_,"a_testf":O_,
 "o_prio":O_,"f_relsch":O_,"o_visit":O_,"o_rate":O_}

gasouthern = {"p_first":R_,"p_last":R_,"p_email":R_,"a_grad":R_,"a_hs":R_,"a_hsaddr":R_,
 "p_photo":O_,"p_middle":O_,"p_hphone":O_,"p_mphone":O_,"p_addr":O_,"p_city":O_,"p_state":O_,"p_zip":O_,"p_country":O_,
 "p_dob":O_,"p_fb":O_,"p_tw":O_,"p_hob":O_,"f_sib":O_,
 "f_p1":O_,"f_coach":O_,"a_couns":O_,"f_infl":O_,
 "a_hsph":O_,"a_trans":O_,"a_testf":O_,"a_gpa":O_,"a_rank":O_,"a_major":O_,"a_ncaa":O_,"a_ncaaid":O_,
 "a_sat":O_,"a_act":O_,"a_hon":O_,"a_other":O_,
 "t_ppos":O_,"t_ht":O_,"t_wt":O_,"t_foot":O_,"t_vid":O_,"t_club":O_,"t_cpos":O_,"t_tourn":O_,"t_osport":O_,"t_stats":O_}

# ---- JumpForward ----
usf = {"p_first":O_,"p_middle":O_,"p_last":O_,"p_pref":O_,"p_country":O_,"p_addr":O_,"p_city":O_,"p_state":O_,"p_zip":O_,
 "p_dob":O_,"p_email":O_,"p_mphone":O_,"p_tw":O_,"p_ig":O_,"a_grad":O_,
 "a_hs":O_,"a_hsaddr":O_,"a_gpa":O_,"a_sat":O_,"a_act":O_,"a_toefl":O_,"a_major":O_,"a_ncaa":O_,"a_ncaaid":O_,
 "t_ht":O_,"t_wt":O_,"t_jersey":O_,"t_ppos":O_,"t_club":O_,"t_cloc":O_,"t_ccoach":O_,"t_tourn":O_,
 "f_p1":O_,"f_p2":O_,"f_pocc":O_,"f_pemail":O_,"f_sib":O_,"f_infl":O_}
scarolina = {"p_first":R_,"p_middle":O_,"p_last":R_,"p_pref":O_,"p_country":R_,"p_addr":R_,"p_city":R_,"p_state":R_,"p_zip":R_,
 "p_dob":R_,"p_email":R_,"p_mphone":R_,
 "a_hs":R_,"a_hsaddr":R_,"a_gpa":O_,"a_sat":O_,"a_act":O_,"a_toefl":O_,"a_major":O_,"a_ncaa":O_,"a_ncaaid":O_,
 "t_ht":R_,"t_wt":R_,"t_jersey":O_,"t_hudl":O_,"t_ppos":R_,"t_tourn":O_,"t_club":R_,"t_cloc":R_,"t_ccoach":R_,
 "o_free":O_}
fiu = {"p_first":O_,"p_middle":O_,"p_last":O_,"p_pref":O_,"p_country":O_,"p_addr":O_,"p_city":O_,"p_state":O_,"p_zip":O_,
 "a_grad":O_,"p_dob":O_,"p_email":O_,"p_hphone":O_,"p_mphone":O_,"p_fb":O_,"p_tw":O_,"p_ig":O_,
 "a_hs":O_,"a_hsaddr":O_,"a_couns":O_,"a_gpa":O_,"a_sat":O_,"a_act":O_,"a_toefl":O_,"a_rank":O_,"a_major":O_,"a_ncaa":O_,"a_ncaaid":O_,
 "t_ht":O_,"t_wt":O_,"t_jersey":O_,"t_club":O_,"t_cloc":O_,"t_ccoach":O_,
 "f_p1":O_,"f_p2":O_,"f_pocc":O_,"f_pemail":O_,"f_sib":O_,"f_infl":O_}

# ---- SIDEARM prospect forms ----
toronto = {"p_full":R_,"p_addr":O_,"p_city":O_,"p_state":O_,"p_zip":O_,"p_country":O_,"p_hphone":O_,"p_mphone":O_,
 "p_email":O_,"p_dob":O_,"f_p1":O_,"o_usid":O_,
 "a_hs":O_,"a_colprev":O_,"a_hsph":O_,"a_grad":O_,"a_major":O_,"a_gpa":O_,"a_sat":O_,"a_act":O_,"a_toefl":O_,"a_hon":O_,"a_other":O_,
 "t_sport":O_,"t_ht":O_,"t_wt":O_,"t_club":O_,"t_ccoach":O_,"t_ppos":O_,"t_stats":O_,"o_bio":O_}
york = {"p_full":R_,"p_addr":O_,"p_city":O_,"p_state":O_,"p_zip":O_,"p_country":O_,"p_hphone":O_,"p_mphone":O_,
 "p_email":R_,"p_dob":O_,"f_p1":O_,
 "a_hs":O_,"a_colprev":O_,"a_hsaddr":O_,"a_hsph":O_,"a_grad":O_,"a_major":O_,"a_gpa":O_,"o_finaid":O_,"a_hon":O_,"a_other":O_,
 "t_sport":O_,"t_osport":O_,"t_ht":O_,"t_wt":O_,"t_club":O_,"t_stats":O_,"t_ccoach":O_,"t_ppos":O_,"o_bio":O_}

# ---- UCF custom bio ----
ucf = {"p_full":R_,"p_pref":O_,"p_pron":O_,"t_ht":O_,"p_home":R_,"p_dob":R_,"p_bplace":R_,
 "a_major":R_,"a_hon":R_,"a_incls":O_,"t_ppos":R_,"a_hs":R_,"a_grad":R_,"t_stats":R_,"t_osport":R_,"a_colprev":O_,
 "f_sib":O_,"f_relath":O_,"f_relsch":O_,"o_why":R_,"p_ig":R_,"p_tw":R_,"o_bio":O_,"o_vol":O_}

# ---- Ohio State FieldLevel (generic profile) ----
ohio = {"p_full":O_,"p_email":O_,"a_grad":O_,"a_gpa":O_,"a_sat":O_,"a_act":O_,"t_ht":O_,"t_wt":O_,
 "t_ppos":O_,"t_club":O_,"a_hs":O_,"t_vid":O_}

SCHOOLS = [
 ("UCF", ucf),("USF", usf),("S. Carolina", scarolina),("Ohio St*", ohio),
 ("Penn St†", pennstate),("Boston U†", bostonu),("Northeastern", northeastern),
 ("UMass Lowell", umasslowell),("IU Indy", iuindy),("FIU", fiu),
 ("NC State", ncstate),("Ga Southern", gasouthern),("Toronto", toronto),("York", york),
]

# header
mx.cell(1,1,"Fields required by each questionnaire — with Mathew's answers").font = Font(bold=True, size=13, color="1F4E79")
mx.merge_cells(start_row=1,start_column=1,end_row=1,end_column=3+len(SCHOOLS))
legend = ("Legend:  R = required   ● = asked (optional / required-status unknown)   blank = not on that form   "
          "|  * Ohio St = FieldLevel account required, generic profile fields   † Penn St / Boston U = ARMS standard template, not individually verified")
mx.cell(2,1,legend).font = Font(italic=True, size=9)
mx.merge_cells(start_row=2,start_column=1,end_row=2,end_column=3+len(SCHOOLS))
mx.cell(2,1).alignment = WRAP

hrow = 3
heads = ["Group","Data field","Mathew's answer"] + [s[0] for s in SCHOOLS]
for j,h in enumerate(heads,1):
    c = mx.cell(hrow,j,h); c.font=HDR; c.fill=HDR_FILL; c.border=BORDER
    c.alignment = Alignment(wrap_text=True, vertical="bottom", horizontal="center" if j>3 else "left")

cur_group = ""
rr = hrow+1
for key,label,ansk in ROWS:
    if label is None:  # group header
        cur_group = key
        c = mx.cell(rr,1,key); c.font=GRP_FONT; c.fill=GRP_FILL
        for j in range(1,4+len(SCHOOLS)):
            mx.cell(rr,j).fill=GRP_FILL; mx.cell(rr,j).border=BORDER
        rr+=1
        continue
    mx.cell(rr,1,cur_group).alignment=WRAP
    mx.cell(rr,2,label).font=Font(bold=True); mx.cell(rr,2).alignment=WRAP
    ans = A.get(ansk,"")
    ac = mx.cell(rr,3,ans); ac.alignment=WRAP
    if isinstance(ans,str) and NEED in ans: ac.fill=NEED_FILL
    else: ac.fill=ANS_FILL
    for j,(name,mp) in enumerate(SCHOOLS, start=4):
        mark = mp.get(key,"")
        c = mx.cell(rr,j,mark); c.alignment=CTR; c.border=BORDER
        if mark==R_: c.font=REQ_FONT
    for j in range(1,4+len(SCHOOLS)):
        mx.cell(rr,j).border=BORDER
    rr+=1

mx.column_dimensions["A"].width = 11
mx.column_dimensions["B"].width = 30
mx.column_dimensions["C"].width = 44
for j in range(4,4+len(SCHOOLS)):
    mx.column_dimensions[get_column_letter(j)].width = 12
mx.freeze_panes = "D4"

out = "/Users/brendonpaul/.immediacy/worktrees/487059d6-f55c-4282-9bd3-f229c01cfb44/Mathew/Mathew-questionnaire-responses.xlsx"
wb.save(out)
print("saved", out)
print("sheets:", wb.sheetnames)
